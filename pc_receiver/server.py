from __future__ import annotations

import argparse
import csv
import io
import json
import os
import shutil
import sqlite3
import sys
import threading
import time
from datetime import datetime, timedelta
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

APP_VERSION = "2.1.0"
DEFAULT_ROOT = r"\\192.168.16.211\Data\Production4\MTS_Result"
SUPERVISORS = {"wichan", "somchai", "supat", "nittaya"}
PLAN_GROUPS = [
    "Cut Rack Bar", "Cut Slug Part", "Chamfer", "Hand Chamfer",
    "Chamfer Slugnut", "Cutpress", "Lathe Bar", "Bending", "Other",
]
MACHINES = {
    **{x: ("Cutting 1", "Cut Rack Bar" if x in {"SC21","SC22","SC23","SC24"} else "Cut Slug Part") for x in ["SC12","SC16","SC25","SC26","SC27","SC28"]},
    **{x: ("Cutting 2", "Cut Rack Bar" if x in {"SC21","SC22","SC23","SC24"} else "Cut Slug Part") for x in ["SC13","SC15","SC17","SC18","SC19","SC20","SC21","SC22","SC23","SC24","CP2","M069"]},
    **{x: ("Chamfer Slugnut 3MC", "Chamfer Slugnut") for x in ["CH5","CH6","CH8"]},
    **{x: ("Chamfer / Hand Chamfer", "Chamfer") for x in ["CH10","CH11","CH2","CH4","CH7"]},
    "HAND_CHAMFER": ("Chamfer / Hand Chamfer", "Hand Chamfer"),
    "BENDING": ("Bending", "Bending"),
    "SCREENING": ("Other", "Other"), "CHECK_RUN_OUT": ("Other", "Other"), "REPAIR": ("Other", "Other"),
}


def now_ms() -> int:
    return int(time.time() * 1000)


def production_date(event_ms: int, shift_name: str) -> str:
    dt = datetime.fromtimestamp(event_ms / 1000)
    if shift_name.upper() == "NIGHT" and dt.hour < 12:
        dt -= timedelta(days=1)
    return dt.strftime("%Y-%m-%d")


class Database:
    def __init__(self, root: Path):
        self.root = root
        self.db_dir = root / "Database"
        self.plan_dir = root / "Import" / "Plan"
        self.result_dir = root / "Import" / "Result"
        self.export_daily = root / "Export" / "Daily"
        self.export_monthly = root / "Export" / "Monthly"
        self.backup_dir = root / "Backup"
        self.log_dir = root / "Logs"
        for p in (self.db_dir, self.plan_dir, self.result_dir, self.export_daily,
                  self.export_monthly, self.backup_dir, self.log_dir):
            p.mkdir(parents=True, exist_ok=True)
        self.path = self.db_dir / "mts_result.db"
        self.lock = threading.RLock()
        self.init_schema()

    def connect(self):
        c = sqlite3.connect(self.path, timeout=30)
        c.row_factory = sqlite3.Row
        c.execute("PRAGMA journal_mode=WAL")
        c.execute("PRAGMA busy_timeout=30000")
        return c

    def init_schema(self):
        with self.connect() as c:
            c.executescript(
                """
                CREATE TABLE IF NOT EXISTS events(
                  event_id TEXT PRIMARY KEY,event_type TEXT NOT NULL,machine_id TEXT NOT NULL,
                  group_name TEXT,plan_group TEXT,employee TEXT,shift_name TEXT,
                  production_date TEXT,event_ms INTEGER,payload_json TEXT,received_ms INTEGER);
                CREATE INDEX IF NOT EXISTS idx_events_period ON events(production_date,shift_name,event_type);
                CREATE INDEX IF NOT EXISTS idx_events_group ON events(plan_group,event_type);
                CREATE TABLE IF NOT EXISTS machine_status(
                  machine_id TEXT PRIMARY KEY,group_name TEXT,plan_group TEXT,status TEXT,
                  employee TEXT,shift_name TEXT,production_date TEXT,item TEXT,part_no TEXT,lot TEXT,
                  reason TEXT,detail TEXT,ok_qty INTEGER DEFAULT 0,ng_qty INTEGER DEFAULT 0,
                  stop_sec INTEGER DEFAULT 0,updated_ms INTEGER);
                CREATE TABLE IF NOT EXISTS approvals(
                  production_date TEXT,shift_name TEXT,supervisor TEXT,approved_ms INTEGER,note TEXT,
                  PRIMARY KEY(production_date,shift_name));
                CREATE TABLE IF NOT EXISTS plan_month(
                  month_key TEXT,group_name TEXT,plan_qty REAL,plan_hours REAL,workdays INTEGER,
                  source_file TEXT,imported_ms INTEGER,PRIMARY KEY(month_key,group_name));
                CREATE TABLE IF NOT EXISTS shift_attendance(
                  production_date TEXT,shift_name TEXT,total INTEGER,sick INTEGER,personal INTEGER,
                  vacation INTEGER,outside INTEGER,outside_reason TEXT,available INTEGER,updated_ms INTEGER,
                  PRIMARY KEY(production_date,shift_name));
                """
            )
            for machine, (group, plan_group) in MACHINES.items():
                c.execute(
                    "INSERT OR IGNORE INTO machine_status(machine_id,group_name,plan_group,status,updated_ms) VALUES(?,?,?,?,0)",
                    (machine, group, plan_group, "UNKNOWN"),
                )

    def ingest(self, event: dict) -> bool:
        required = ("event_id", "event_type", "machine_id", "event_ms")
        if any(not event.get(k) for k in required):
            raise ValueError("missing required event field")
        payload = event.get("payload") or {}
        if not isinstance(payload, dict):
            raise ValueError("payload must be an object")
        shift = str(event.get("shift_name") or "")
        pdate = production_date(int(event["event_ms"]), shift)
        values = (
            str(event["event_id"]), str(event["event_type"]), str(event["machine_id"]),
            str(event.get("group_name") or ""), str(event.get("plan_group") or ""),
            str(event.get("employee") or ""), shift, pdate, int(event["event_ms"]),
            json.dumps(payload, ensure_ascii=False), now_ms(),
        )
        with self.lock, self.connect() as c:
            cur = c.execute(
                "INSERT OR IGNORE INTO events VALUES(?,?,?,?,?,?,?,?,?,?,?)", values
            )
            if cur.rowcount == 0:
                return False
            self._update_status(c, event, payload, pdate)
            return True

    def _update_status(self, c, e: dict, p: dict, pdate: str):
        if str(e["event_type"]) == "ATTENDANCE":
            c.execute(
                """INSERT OR REPLACE INTO shift_attendance
                   (production_date,shift_name,total,sick,personal,vacation,outside,outside_reason,available,updated_ms)
                   VALUES(?,?,?,?,?,?,?,?,?,?)""",
                (pdate, str(e.get("shift_name") or ""), int(p.get("total", 0) or 0),
                 int(p.get("sick", 0) or 0), int(p.get("personal", 0) or 0),
                 int(p.get("vacation", 0) or 0), int(p.get("outside", 0) or 0),
                 str(p.get("outside_reason") or ""), int(p.get("available", 0) or 0),
                 int(e["event_ms"])),
            )
            return
        machine = str(e["machine_id"])
        old = c.execute("SELECT * FROM machine_status WHERE machine_id=?", (machine,)).fetchone()
        row = dict(old) if old else {
            "machine_id": machine, "group_name": e.get("group_name", ""),
            "plan_group": e.get("plan_group", ""), "status": "UNKNOWN",
            "employee": "", "shift_name": "", "production_date": pdate,
            "item": "", "part_no": "", "lot": "", "reason": "", "detail": "",
            "ok_qty": 0, "ng_qty": 0, "stop_sec": 0, "updated_ms": 0,
        }
        row.update({"group_name": e.get("group_name", row["group_name"]),
                    "plan_group": e.get("plan_group", row["plan_group"]),
                    "employee": e.get("employee", ""), "shift_name": e.get("shift_name", ""),
                    "production_date": pdate, "updated_ms": int(e["event_ms"])})
        et = str(e["event_type"])
        if et == "STATUS":
            row["status"] = p.get("status", row["status"])
            row["reason"] = p.get("reason", "")
            row["detail"] = p.get("detail", "")
            if row["status"] == "RUNNING" and old and old["shift_name"] != row["shift_name"]:
                row["ok_qty"] = row["ng_qty"] = row["stop_sec"] = 0
        elif et == "TAG":
            row["status"] = "RUNNING"
            row["item"] = p.get("item", "")
            row["part_no"] = p.get("part_no", "")
            row["lot"] = p.get("lot", "")
            row["ok_qty"] = int(row.get("ok_qty", 0)) + int(p.get("this_qty", 0) or 0)
        elif et == "NG":
            row["ng_qty"] = int(row.get("ng_qty", 0)) + int(p.get("qty", 0) or 0)
        elif et == "CLOSE_SHIFT":
            row["status"] = "CLOSED"
            row["ok_qty"] = int(p.get("total_ok", row.get("ok_qty", 0)) or 0)
            row["ng_qty"] = int(p.get("total_ng", row.get("ng_qty", 0)) or 0)
            row["stop_sec"] = int(p.get("stop_sec", row.get("stop_sec", 0)) or 0)
        cols = list(row.keys())
        c.execute(
            f"INSERT OR REPLACE INTO machine_status({','.join(cols)}) VALUES({','.join('?' for _ in cols)})",
            tuple(row[x] for x in cols),
        )

    def approve(self, day: str, shift: str, supervisor: str, note: str = ""):
        name = supervisor.strip()
        if name.lower() not in SUPERVISORS:
            raise ValueError("ผู้ใช้นี้ไม่มีสิทธิ์ยืนยัน")
        datetime.strptime(day, "%Y-%m-%d")
        shift = shift.upper()
        if shift not in {"DAY", "NIGHT"}:
            raise ValueError("shift must be DAY or NIGHT")
        with self.connect() as c:
            closed = c.execute(
                "SELECT COUNT(*) FROM events WHERE production_date=? AND shift_name=? AND event_type='CLOSE_SHIFT'",
                (day, shift),
            ).fetchone()[0]
            if not closed:
                raise ValueError("ยังไม่มีเครื่องปิดกะในวันที่และกะนี้")
            reported = {r[0] for r in c.execute(
                "SELECT DISTINCT machine_id FROM events WHERE production_date=? AND shift_name=?",
                (day, shift),
            )}
            missing = [x for x in MACHINES if x not in reported]
            if missing:
                raise ValueError("ยังไม่มีสถานะ: " + ", ".join(missing))
            c.execute(
                "INSERT OR REPLACE INTO approvals VALUES(?,?,?,?,?)",
                (day, shift, name, now_ms(), note.strip()),
            )

    def status_rows(self):
        with self.connect() as c:
            return [dict(r) for r in c.execute("SELECT * FROM machine_status ORDER BY group_name,machine_id")]

    def approval_rows(self, limit=100):
        with self.connect() as c:
            return [dict(r) for r in c.execute(
                "SELECT * FROM approvals ORDER BY production_date DESC,shift_name LIMIT ?", (limit,)
            )]

    def dashboard(self, month: str | None = None):
        month = month or datetime.now().strftime("%Y-%m")
        with self.connect() as c:
            plans = {r["group_name"]: dict(r) for r in c.execute(
                "SELECT * FROM plan_month WHERE month_key=?", (month,)
            )}
            qty = {r["plan_group"]: r["qty"] for r in c.execute(
                """SELECT e.plan_group,SUM(CAST(json_extract(e.payload_json,'$.this_qty') AS INTEGER)) qty
                   FROM events e JOIN approvals a ON a.production_date=e.production_date AND a.shift_name=e.shift_name
                   WHERE substr(e.production_date,1,7)=? AND e.event_type='TAG' GROUP BY e.plan_group""", (month,)
            )}
            hours = {r["plan_group"]: r["hrs"] for r in c.execute(
                """SELECT e.plan_group,SUM(CAST(json_extract(e.payload_json,'$.working_sec') AS REAL))/3600.0 hrs
                   FROM events e JOIN approvals a ON a.production_date=e.production_date AND a.shift_name=e.shift_name
                   WHERE substr(e.production_date,1,7)=? AND e.event_type='CLOSE_SHIFT' GROUP BY e.plan_group""", (month,)
            )}
            days = c.execute("SELECT COUNT(DISTINCT production_date) FROM approvals WHERE substr(production_date,1,7)=?", (month,)).fetchone()[0]
            daily_qty = {r["production_date"]: float(r["qty"] or 0) for r in c.execute(
                """SELECT e.production_date,SUM(CAST(json_extract(e.payload_json,'$.this_qty') AS INTEGER)) qty
                   FROM events e JOIN approvals a ON a.production_date=e.production_date AND a.shift_name=e.shift_name
                   WHERE substr(e.production_date,1,7)=? AND e.event_type='TAG' GROUP BY e.production_date ORDER BY e.production_date""", (month,)
            )}
            approved_events = [dict(r) for r in c.execute(
                """SELECT e.machine_id,e.production_date,e.shift_name,e.event_type,e.event_ms,e.payload_json
                   FROM events e JOIN approvals a ON a.production_date=e.production_date AND a.shift_name=e.shift_name
                   WHERE substr(e.production_date,1,7)=? ORDER BY e.machine_id,e.production_date,e.shift_name,e.event_ms""", (month,)
            )]
        groups = []
        all_groups = list(dict.fromkeys(PLAN_GROUPS + list(plans) + list(qty) + list(hours)))
        total_plan = total_actual = total_plan_hours = total_actual_hours = 0.0
        workdays = max([int(x.get("workdays") or 0) for x in plans.values()] or [0])
        progress = (days / workdays) if workdays else None
        for group in all_groups:
            p = float(plans.get(group, {}).get("plan_qty") or 0)
            ph = float(plans.get(group, {}).get("plan_hours") or 0)
            a = float(qty.get(group, 0) or 0)
            ah = float(hours.get(group, 0) or 0)
            month_pct = a / p if p else None
            to_date_pct = a / (p * progress) if p and progress else None
            groups.append({"group": group, "plan_qty": p, "actual_qty": a,
                           "month_pct": month_pct, "to_date_pct": to_date_pct,
                           "plan_hours": ph, "actual_hours": ah,
                           "hour_pct": ah / ph if ph else None,
                           "status": "NO_PLAN" if not p else ("OK" if to_date_pct is not None and to_date_pct >= 1 else "BEHIND")})
            if p:
                total_plan += p; total_actual += a; total_plan_hours += ph; total_actual_hours += ah
        daily_target = total_plan / workdays if workdays else 0
        daily = [{"date": d, "plan_qty": daily_target, "actual_qty": a,
                  "pct": a / daily_target if daily_target else None} for d, a in sorted(daily_qty.items())]
        losses = self._loss_summary(approved_events)
        return {"month": month, "days_passed": days, "workdays": workdays,
                "time_progress": progress, "total_plan": total_plan, "total_actual": total_actual,
                "total_month_pct": total_actual / total_plan if total_plan else None,
                "total_to_date_pct": total_actual / (total_plan * progress) if total_plan and progress else None,
                "total_plan_hours": total_plan_hours, "total_actual_hours": total_actual_hours,
                "groups": groups, "daily": daily, "losses": losses}

    @staticmethod
    def _loss_summary(events):
        totals = {}
        active = {}
        for e in events:
            key = (e["machine_id"], e["production_date"], e["shift_name"])
            payload = json.loads(e["payload_json"] or "{}")
            et = e["event_type"]
            if key in active and (et in {"STATUS", "TAG", "CLOSE_SHIFT"}):
                start, reason = active.pop(key)
                totals[reason] = totals.get(reason, 0) + max(0, int(e["event_ms"]) - start)
            if et == "STATUS" and payload.get("status") in {"PLANNED_STOP", "UNPLANNED_STOP", "SETUP"}:
                reason = str(payload.get("reason") or payload.get("status") or "ไม่ระบุ")
                detail = str(payload.get("detail") or "").strip()
                label = reason + ((" — " + detail) if detail else "")
                active[key] = (int(e["event_ms"]), label)
        return [{"reason": k, "hours": v / 3600000.0} for k, v in sorted(totals.items(), key=lambda x: x[1], reverse=True)]

    def import_plan(self, path: Path):
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb["เปรียบเทียบ"] if "เปรียบเทียบ" in wb.sheetnames else wb[wb.sheetnames[0]]
        source_date = ws["B1"].value
        if isinstance(source_date, datetime):
            month = source_date.strftime("%Y-%m")
        else:
            month = datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m")
        workdays = int(ws["D15"].value or 0)
        mapping = {"Run-Out": "Other", "Screening": "Other", "Repair": "Other"}
        combined = {}
        for col in range(2, 15):
            group = str(ws.cell(4, col).value or "").strip()
            if not group or group == "Total (Report)":
                continue
            group = mapping.get(group, group)
            plan_qty = float(ws.cell(5, col).value or 0)
            plan_hours = float(ws.cell(10, col).value or 0)
            old = combined.get(group, (0.0, 0.0))
            combined[group] = (old[0] + plan_qty, old[1] + plan_hours)
        rows = [(month, group, values[0], values[1], workdays, path.name, now_ms()) for group, values in combined.items()]
        with self.connect() as c:
            c.executemany("INSERT OR REPLACE INTO plan_month VALUES(?,?,?,?,?,?,?)", rows)
        return {"month": month, "groups": len(rows), "workdays": workdays}

    def import_latest_plan(self):
        files = sorted(self.plan_dir.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
        if not files:
            return None
        return self.import_plan(files[0])

    def export_xlsx(self, month: str | None = None) -> Path:
        from openpyxl import Workbook
        month = month or datetime.now().strftime("%Y-%m")
        data = self.dashboard(month)
        wb = Workbook(); ws = wb.active; ws.title = "Dashboard"
        ws.append(["Month", month, "Days Passed", data["days_passed"], "Workdays", data["workdays"]])
        ws.append(["Group", "Plan Qty", "Actual Qty", "% Month", "% To Date", "Plan Hr", "Actual Hr", "Status"])
        for r in data["groups"]:
            ws.append([r["group"], r["plan_qty"], r["actual_qty"], r["month_pct"], r["to_date_pct"], r["plan_hours"], r["actual_hours"], r["status"]])
        for cell in ws[2]: cell.font = __import__("openpyxl").styles.Font(bold=True)
        for row in range(3, ws.max_row + 1): ws.cell(row, 4).number_format = ws.cell(row, 5).number_format = "0.0%"
        dy = wb.create_sheet("Daily"); dy.append(["Production Date","Daily Plan Qty","Actual Qty","% Daily Plan"])
        for r in data["daily"]: dy.append([r["date"],r["plan_qty"],r["actual_qty"],r["pct"]])
        for row in range(2,dy.max_row+1): dy.cell(row,4).number_format="0.0%"
        ls = wb.create_sheet("Losses"); ls.append(["Stop Reason / Detail","Lost Hours"])
        for r in data["losses"]: ls.append([r["reason"],r["hours"]])
        ev = wb.create_sheet("Events"); ev.append(["Date","Shift","Type","Machine","Group","Plan Group","Employee","Time","Payload"])
        with self.connect() as c:
            for r in c.execute("SELECT production_date,shift_name,event_type,machine_id,group_name,plan_group,employee,event_ms,payload_json FROM events WHERE substr(production_date,1,7)=? ORDER BY event_ms", (month,)):
                ev.append([r[0],r[1],r[2],r[3],r[4],r[5],r[6],datetime.fromtimestamp(r[7]/1000),r[8]])
        ap = wb.create_sheet("Approvals"); ap.append(["Production Date","Shift","Supervisor","Approved Time","Note"])
        for r in self.approval_rows(10000): ap.append([r["production_date"],r["shift_name"],r["supervisor"],datetime.fromtimestamp(r["approved_ms"]/1000),r["note"]])
        at = wb.create_sheet("Attendance"); at.append(["Production Date","Shift","Total","Sick","Personal","Vacation","Outside","Outside Reason","Available","Updated"])
        with self.connect() as c:
            for r in c.execute("SELECT * FROM shift_attendance WHERE substr(production_date,1,7)=? ORDER BY production_date,shift_name", (month,)):
                at.append([r["production_date"],r["shift_name"],r["total"],r["sick"],r["personal"],r["vacation"],r["outside"],r["outside_reason"],r["available"],datetime.fromtimestamp(r["updated_ms"]/1000)])
        st = wb.create_sheet("Machine Status"); rows=self.status_rows(); st.append(list(rows[0].keys()) if rows else ["No data"])
        for r in rows: st.append(list(r.values()))
        out = self.export_monthly / f"MTS_Result_{month}.xlsx"; wb.save(out); return out

    def backup(self):
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        target = self.backup_dir / f"mts_result_{stamp}.db"
        with self.lock, self.connect() as src, sqlite3.connect(target) as dst:
            src.backup(dst)
        return target


class App:
    def __init__(self, db: Database, host: str, port: int):
        self.db, self.host, self.port = db, host, port
        self.dashboard_html = (Path(__file__).with_name("dashboard.html")).read_text(encoding="utf-8")

    def handler(self):
        app = self
        class Handler(BaseHTTPRequestHandler):
            server_version = f"MTSResult/{APP_VERSION}"
            def log_message(self, fmt, *args):
                line = f"{datetime.now().isoformat(timespec='seconds')} {self.client_address[0]} {fmt % args}\n"
                with (app.db.log_dir / "receiver.log").open("a", encoding="utf-8") as f: f.write(line)
            def send_json(self, obj, status=200):
                raw=json.dumps(obj,ensure_ascii=False).encode("utf-8");self.send_response(status);self.send_header("Content-Type","application/json; charset=utf-8");self.send_header("Content-Length",str(len(raw)));self.end_headers();self.wfile.write(raw)
            def body_json(self):
                n=int(self.headers.get("Content-Length","0"));return json.loads(self.rfile.read(n).decode("utf-8"))
            def do_GET(self):
                u=urlparse(self.path);q=parse_qs(u.query)
                try:
                    if u.path=="/":
                        raw=app.dashboard_html.encode("utf-8");self.send_response(200);self.send_header("Content-Type","text/html; charset=utf-8");self.send_header("Content-Length",str(len(raw)));self.end_headers();self.wfile.write(raw)
                    elif u.path=="/api/health": self.send_json({"ok":True,"version":APP_VERSION,"data_root":str(app.db.root)})
                    elif u.path=="/api/status": self.send_json({"rows":app.db.status_rows()})
                    elif u.path=="/api/dashboard": self.send_json(app.db.dashboard(q.get("month",[None])[0]))
                    elif u.path=="/api/approvals": self.send_json({"rows":app.db.approval_rows()})
                    elif u.path=="/export.xlsx":
                        p=app.db.export_xlsx(q.get("month",[None])[0]);raw=p.read_bytes();self.send_response(200);self.send_header("Content-Type","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet");self.send_header("Content-Disposition",f'attachment; filename="{p.name}"');self.send_header("Content-Length",str(len(raw)));self.end_headers();self.wfile.write(raw)
                    else:self.send_json({"error":"not found"},404)
                except Exception as e:self.send_json({"error":str(e)},500)
            def do_POST(self):
                try:
                    if self.path=="/api/events": self.send_json({"accepted":app.db.ingest(self.body_json())})
                    elif self.path=="/api/approve":
                        b=self.body_json();app.db.approve(str(b.get("production_date","")),str(b.get("shift_name","")),str(b.get("supervisor","")),str(b.get("note","")));self.send_json({"ok":True})
                    elif self.path=="/api/import-plan": self.send_json({"ok":True,"result":app.db.import_latest_plan()})
                    elif self.path=="/api/backup": self.send_json({"ok":True,"file":str(app.db.backup())})
                    else:self.send_json({"error":"not found"},404)
                except (ValueError,KeyError,json.JSONDecodeError) as e:self.send_json({"error":str(e)},400)
                except Exception as e:self.send_json({"error":str(e)},500)
        return Handler

    def run(self):
        server=ThreadingHTTPServer((self.host,self.port),self.handler())
        print(f"MTS Result Receiver v{APP_VERSION}")
        print(f"Dashboard: http://{self.host if self.host!='0.0.0.0' else '192.168.18.145'}:{self.port}")
        print(f"Data: {self.db.root}")
        server.serve_forever()


def main(argv=None):
    p=argparse.ArgumentParser();p.add_argument("--data-root",default=os.environ.get("MTS_DATA_ROOT",DEFAULT_ROOT));p.add_argument("--host",default="0.0.0.0");p.add_argument("--port",type=int,default=8765);p.add_argument("--import-plan",action="store_true");p.add_argument("--export",action="store_true");args=p.parse_args(argv)
    db=Database(Path(args.data_root))
    if args.import_plan: print(db.import_latest_plan());return 0
    if args.export: print(db.export_xlsx());return 0
    App(db,args.host,args.port).run();return 0


if __name__=="__main__": raise SystemExit(main())
